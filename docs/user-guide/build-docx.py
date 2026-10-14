#!/usr/bin/env python3
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "Unet-LMS-rukovodstvo-prepodavatelya.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "Calibri", 18 if level == 1 else 14 if level == 2 else 12, True)
    return p


def add_p(doc, text, size=11, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, "Calibri", size, bold, color)
    run.italic = italic
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run, "Calibri", 11)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        for run in p.runs:
            set_run_font(run, "Calibri", 11)


def add_figure(doc, filename, caption):
    path = SHOTS / filename
    if not path.exists():
        add_p(doc, f"[Нет скриншота: {filename}]", italic=True, color=(120, 120, 120))
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, "Calibri", 9, False, (90, 90, 90))
    cap.paragraph_format.space_after = Pt(12)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, "Calibri", 10, True)
    for r_i, row in enumerate(rows):
        for c_i, value in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, "Calibri", 10)
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = doc.add_paragraph()
    t = title.add_run("UNET LMS")
    set_run_font(t, "Calibri", 12, True, (90, 90, 90))
    title.paragraph_format.space_after = Pt(4)

    h = doc.add_paragraph()
    run = h.add_run("Руководство пользователя\nдля преподавателя")
    set_run_font(run, "Calibri", 26, True)
    h.paragraph_format.space_after = Pt(10)

    add_p(
        doc,
        "Актуальный интерфейс: вход, страница «Сегодня», курсы, темы, материалы, приглашения, тесты и коллекция вопросов.",
        size=12,
        color=(70, 70, 70),
    )
    add_p(
        doc,
        "Роль: преподаватель (профиль сотрудника UNET). https://lms.kstu.kg · сентябрь 2026.",
        size=10,
        color=(110, 110, 110),
    )

    add_heading(doc, "Содержание", 1)
    add_bullets(
        doc,
        [
            "1. Вход в систему",
            "2. Страница «Сегодня»",
            "3. Раздел «Мои курсы»",
            "4. Как добавить курс",
            "5. Страница курса",
            "6. Как добавить тему или тест в курс",
            "7. Работа с выбранной темой",
            "8. Учебные материалы",
            "9. Приглашение студентов",
            "10. Успеваемость и управление курсом",
            "11. Как создать тест",
            "12. Коллекция вопросов",
        ],
    )

    add_heading(doc, "1. Вход в систему", 1)
    add_p(
        doc,
        "Unet LMS использует те же учётные данные, что и корпоративная система UNET. "
        "Входите логином (ПИН) и паролем UNET. Если у учётной записи несколько профилей, "
        "для работы преподавателя выберите «Сотрудник».",
    )
    add_figure(doc, "01-login.png", "Рисунок 1. Экран входа: ПИН, пароль, вход через корпоративную почту")
    add_heading(doc, "Что заполнить", 2)
    add_table(
        doc,
        ["Поле", "Что вводить"],
        [
            ["ПИН", "Логин из системы UNET"],
            ["Пароль", "Пароль от UNET. По умолчанию у сотрудника — ИНН, у студента — «s» + ИНН"],
        ],
    )
    add_steps(
        doc,
        [
            "Откройте https://lms.kstu.kg.",
            "Введите ПИН и пароль.",
            "Нажмите «Войти».",
            "Если появится выбор профиля, выберите «Сотрудник». Не выбирайте «Студент».",
        ],
    )
    add_p(
        doc,
        "Дополнительно доступен вход через «Корпоративная почта». После входа открывается страница «Сегодня». "
        "В шапке: Сегодня, Мои курсы, Тестирование, Коллекция вопросов, Нагрузка (если есть право), Руководство.",
        italic=True,
    )

    add_heading(doc, "2. Страница «Сегодня»", 1)
    add_p(
        doc,
        "Стартовый экран преподавателя. Слева очередь «На проверку» (сначала те, кто сдал раньше). "
        "Справа — фильтр по курсам, календарь сроков и объявления. Точка в календаре — день открытия задания.",
    )
    add_figure(doc, "02-today.png", "Рисунок 2. «Сегодня»: очередь проверки, календарь и объявления")

    add_heading(doc, "3. Раздел «Мои курсы»", 1)
    add_p(
        doc,
        "Список дисциплин. Вкладки «Курсы» и «Архив». На карточке: студенты, проверено, средний балл, "
        "преподаватель, кнопки «Архив» и «Открыть». Пунктирная карточка создаёт новый курс.",
    )
    add_figure(doc, "03-my-courses.png", "Рисунок 3. Карточки курсов и «Добавить курс»")

    add_heading(doc, "4. Как добавить курс", 1)
    add_steps(
        doc,
        [
            "Откройте «Мои курсы».",
            "Нажмите карточку «Добавить курс».",
            "Заполните форму справа и нажмите «Добавить курс».",
        ],
    )
    add_figure(doc, "04-add-course.png", "Рисунок 4. Панель создания курса: название и кафедра")
    add_table(
        doc,
        ["Поле", "Обязательно", "Описание"],
        [
            ["Название курса", "Да", "До 100 символов."],
            ["Кафедра", "Да", "Организация из профиля сотрудника."],
        ],
    )
    add_p(
        doc,
        "Кредиты и форма контроля при создании курса больше не заполняются. "
        "Название и кафедру можно изменить во вкладке «Управление курсом».",
        italic=True,
    )

    add_heading(doc, "5. Страница курса", 1)
    add_p(doc, "Нажмите «Открыть». Рядом с названием — QR для приглашения. Вкладки:")
    add_bullets(
        doc,
        [
            "Учебный процесс — темы и тесты по типу занятия;",
            "Лента курса — объявления и журнал загрузок;",
            "Материалы — все файлы курса;",
            "О курсе — описание;",
            "Успеваемость студентов — баллы и Excel;",
            "Управление курсом — настройки, доступ, потоки, дублирование, удаление.",
        ],
    )
    add_figure(doc, "05-course-study.png", "Рисунок 5. Учебный процесс")
    add_figure(doc, "10-course-feed.png", "Рисунок 6. Лента курса")

    add_heading(doc, "6. Как добавить тему или тест в курс", 1)
    add_steps(
        doc,
        [
            "Откройте курс, вкладка «Учебный процесс».",
            "Нажмите «Добавить».",
            "Выберите тип занятия и заполните поля.",
        ],
    )
    add_figure(doc, "06-add-theme.png", "Рисунок 7. Панель «Создание темы»")
    add_figure(doc, "07-add-theme-types.png", "Рисунок 8. Типы занятия")
    add_heading(doc, "Оцениваемые типы", 2)
    add_p(
        doc,
        "Лекции, практика, лабораторные, СРС, РГЗ/РП: неделя 1–16, баллы, срок доступа и опция «Закрытый». "
        "Для силлабуса, глоссария, литературы и «другое» баллы скрыты.",
    )
    add_figure(doc, "15-add-theme-lecture.png", "Рисунок 9. Тип «Лекционные материалы»")
    add_heading(doc, "Тип «Тесты»", 2)
    add_p(doc, "Нужно выбрать уже созданный тест из раздела «Тестирование» и привязать его к курсу.")
    add_figure(doc, "16-add-theme-test.png", "Рисунок 10. Выбор готового теста")
    add_p(
        doc,
        "Если теста нет в списке, сначала создайте его в «Тестировании».",
        italic=True,
        color=(146, 64, 14),
    )

    add_heading(doc, "7. Работа с выбранной темой", 1)
    add_p(
        doc,
        "Нажмите тему слева. Справа — материалы, доступ студентов, ответы и обсуждение. "
        "Точечный доступ задаётся здесь, массовый — в «Управлении курсом».",
    )
    add_figure(doc, "08-theme-selected.png", "Рисунок 11. Тема выбрана")

    add_heading(doc, "8. Учебные материалы", 1)
    add_p(
        doc,
        "Нажмите «Добавить материал» или перетащите файл. Допускаются PDF и изображения (лучше PDF). "
        "Файл и ссылка взаимоисключающие. Все файлы курса также во вкладке «Материалы».",
    )
    add_figure(doc, "09-add-material.png", "Рисунок 12. Загрузка материала")
    add_figure(doc, "11-course-files.png", "Рисунок 13. Сводный список файлов")

    add_heading(doc, "9. Приглашение студентов", 1)
    add_p(
        doc,
        "Иконка QR рядом с названием курса. Создайте приглашение и выберите срок. "
        "Чтобы курс увидела учебная группа, добавьте поток в «Управлении курсом».",
    )
    add_figure(doc, "13-invite.png", "Рисунок 14. Диалог приглашения")

    add_heading(doc, "10. Успеваемость и управление курсом", 1)
    add_p(doc, "Успеваемость: баллы, карточка студента, экспорт в Excel.")
    add_figure(doc, "13-students-progress.png", "Рисунок 15. Успеваемость")
    add_p(
        doc,
        "Управление: название и описание, открыть/закрыть доступ ко всем темам, потоки, кафедра, "
        "дублирование (без студентов) и удаление (нельзя, пока есть студенты).",
    )
    add_figure(doc, "14-course-management.png", "Рисунок 16. Управление курсом")

    add_heading(doc, "11. Как создать тест", 1)
    add_p(doc, "Тест создаётся в «Тестировании», затем привязывается к курсу типом «Тесты».")
    add_figure(doc, "14-my-tests.png", "Рисунок 17. Мои тесты")
    add_figure(doc, "15-add-quiz.png", "Рисунок 18. Параметры нового теста")
    add_table(
        doc,
        ["Поле", "Обязательно", "Описание"],
        [
            ["Название теста", "Да", "Как тест виден студентам."],
            ["Описание", "Нет", "Что проверяет тест."],
            ["Время, мин", "Да", "Минимум 1 минута."],
            ["Макс. балл", "Да", "По умолчанию 100."],
            ["Мин. балл для сдачи", "Да", "Не больше максимума."],
            ["Вопросов студенту", "Да", "Случайно из пула, не больше числа вопросов."],
            ["Обязательный", "Нет", "Для выставления балла нужно пройти тест."],
            ["Показывать ответы", "Нет", "После сдачи видны верные ответы."],
        ],
    )
    add_p(
        doc,
        "Типы вопросов: один вариант, несколько, да/нет, короткий ответ, эссе. "
        "Готовые вопросы вставляются кнопкой «Из коллекции вопросов».",
    )
    add_figure(doc, "16-add-quiz-questions.png", "Рисунок 19. Карточка вопроса")
    add_steps(
        doc,
        [
            "Заполните параметры и вопросы.",
            "Укажите, сколько вопросов получит студент.",
            "Нажмите «Создать тест».",
            "В курсе: Добавить → «Тесты» → выберите тест → «Добавить задание».",
        ],
    )

    add_heading(doc, "12. Коллекция вопросов", 1)
    add_p(doc, "Наборы вопросов для вставки в тесты. Создайте коллекцию и наполните её.")
    add_figure(doc, "17-question-bank.png", "Рисунок 20. Коллекции вопросов")

    add_heading(doc, "Краткий порядок работы", 1)
    add_steps(
        doc,
        [
            "Войдите логином UNET, профиль «Сотрудник».",
            "Создайте курс: название и кафедра.",
            "Откройте курс, создайте приглашение или добавьте поток.",
            "Добавьте темы.",
            "При необходимости создайте тест и привяжите его к курсу.",
            "Загрузите материалы в тему.",
            "Проверяйте сдачи на «Сегодня» и смотрите успеваемость.",
        ],
    )

    doc.save(OUT)
    print(f"DOCX: {OUT}")


if __name__ == "__main__":
    main()
